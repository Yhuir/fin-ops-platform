from __future__ import annotations

import unittest
from io import BytesIO
from zipfile import ZipFile
from unittest.mock import MagicMock, Mock, patch

from openpyxl import Workbook

from fin_ops_platform.services.tax_certified_import_application_service import TaxCertifiedImportApplicationService
from fin_ops_platform.services.tax_certified_import_service import (
    SOURCE_COLUMNS,
    TaxCertifiedImportService,
    UploadedCertifiedImportFile,
    _build_unique_key,
)


def certified_upload(*, deductible="0", amount="100", tax="13", month="202609", digital="TEST-DIGITAL-1",
                     buyer="TEST-BUYER", invoice_code="CODE-1", selection_time="2026-10-02 13:14:15", label="增值税专用发票",
                     bad_dimensions=False, extra_rows=None) -> UploadedCertifiedImportFile:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "发票"
    sheet.append(["发票清单"])
    sheet.append(["纳税人识别号", buyer, None, None, "税款所属期", month, None, "纳税人名称", "测试买方"])
    sheet.append(list(SOURCE_COLUMNS))
    row = [1, "已勾选", "电子发票", None, digital, invoice_code, "NUMBER-1", "2026-08-15",
           "TEST-SELLER", "测试销方", amount, tax, deductible, "01", label, "正常", selection_time, "低", "正常"]
    sheet.append(row)
    for extra in extra_rows or []:
        sheet.append(extra)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    if bad_dimensions:
        import re
        source, target = ZipFile(BytesIO(buffer.getvalue())), BytesIO()
        with source, ZipFile(target, "w") as dest:
            for info in source.infolist():
                content = source.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    content = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', content)
                dest.writestr(info, content)
        buffer = target
    return UploadedCertifiedImportFile("用途确认信息.xlsx", buffer.getvalue())


class TaxCertifiedImportServiceTests(unittest.TestCase):
    def preview(self, **kwargs):
        return TaxCertifiedImportService(repository=Mock()).preview_files(imported_by="test-owner", uploads=[certified_upload(**kwargs)]).files[0]

    def test_bad_dimensions_source_labels_zero_and_dates_are_independent(self):
        file = self.preview(bad_dimensions=True)
        self.assertEqual((file.recognized_count, file.invalid_count, file.month), (1, 0, "2026-09"))
        row = file.rows[0]
        self.assertEqual(row.buyer_tax_no, "TEST-BUYER")
        self.assertEqual(row.issue_date, "2026-08-15")
        self.assertEqual(row.selection_time, "2026-10-02 13:14:15")
        self.assertEqual(row.deductible_tax_amount, "0")
        self.assertEqual(row.invoice_kind, "01")
        self.assertEqual(row.invoice_kind_label, "增值税专用发票")
        self.assertEqual(len(row.source_fields), 19)

    def test_missing_amounts_are_not_filled(self):
        row = self.preview(amount=None, tax=None, deductible=None).rows[0]
        self.assertIsNone(row.amount)
        self.assertIsNone(row.tax_amount)
        self.assertIsNone(row.deductible_tax_amount)

    def test_missing_month_does_not_use_selection_date_or_filename(self):
        self.assertIsNone(self.preview(month=None).month)

    def test_non_special_and_invalid_money_and_date_are_explicit(self):
        for kwargs in ({"amount": "NaN"}, {"tax": "bad"}, {"selection_time": "2026-99-99"}):
            with self.subTest(kwargs=kwargs):
                file = self.preview(**kwargs)
                self.assertEqual(file.recognized_count, 0)
                self.assertEqual(file.invalid_count, 1)
                self.assertTrue(file.row_results[0]["error_message"])

    def test_strong_identity_only(self):
        with self.assertRaisesRegex(ValueError, "缺少"):
            _build_unique_key(digital_invoice_no=None, invoice_code=None, invoice_no="123")
        self.assertEqual(_build_unique_key(digital_invoice_no="123", invoice_code="X", invoice_no="Y"), "digital:123")
        self.assertEqual(_build_unique_key(digital_invoice_no=None, invoice_code="X", invoice_no="Y"), "invoice:X:Y")

    def test_empty_uploads_and_conflicting_metadata_rejected(self):
        service = TaxCertifiedImportService(repository=Mock())
        with self.assertRaises(ValueError):
            service.preview_files(imported_by="owner", uploads=[])
        for kwargs in ({"month": "2026-13"}, {"month": "2026-08"}, {"buyer_tax_no": "OTHER"}):
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                service.preview_files(imported_by="owner", uploads=[certified_upload()], **kwargs)

    def test_non_special_is_ignored_and_malformed_xlsx_is_rejected(self):
        file = self.preview(label="普通发票")
        self.assertEqual((file.recognized_count, file.invalid_count, file.ignored_count), (0, 0, 1))
        service = TaxCertifiedImportService(repository=Mock())
        with self.assertRaisesRegex(ValueError, "XLSX"):
            service.preview_files(imported_by="owner", uploads=[UploadedCertifiedImportFile("broken.xlsx", b"not zip")])
        valid = certified_upload()
        buffer = BytesIO()
        with ZipFile(BytesIO(valid.content)) as source, ZipFile(buffer, "w") as destination:
            for info in source.infolist():
                data = b"<broken" if info.filename == "xl/worksheets/sheet1.xml" else source.read(info.filename)
                destination.writestr(info, data)
        with self.assertRaisesRegex(ValueError, "XLSX"):
            service.preview_files(imported_by="owner", uploads=[UploadedCertifiedImportFile("broken.xlsx", buffer.getvalue())])

    def test_parser_stops_reading_at_row_and_column_limits_and_closes_workbook(self):
        workbook = MagicMock()
        workbook.sheetnames = ["发票"]
        sheet = workbook.__getitem__.return_value
        consumed = []
        def rows():
            yield tuple(SOURCE_COLUMNS)
            for number in range(20_005):
                consumed.append(number)
                yield (number, "已勾选")
        sheet.iter_rows.return_value = rows()
        with patch("fin_ops_platform.services.tax_certified_import_service.load_workbook", return_value=workbook):
            with self.assertRaisesRegex(ValueError, "20000"):
                self.preview()
        self.assertEqual(len(consumed), 20_001)
        workbook.close.assert_called_once()
        workbook.close.reset_mock()
        sheet.iter_rows.return_value = iter([tuple(range(65))])
        with patch("fin_ops_platform.services.tax_certified_import_service.load_workbook", return_value=workbook):
            with self.assertRaisesRegex(ValueError, "64"):
                self.preview()
        workbook.close.assert_called_once()

    def test_session_budget_is_shared_across_files_and_empty_rows_are_bounded(self):
        service = TaxCertifiedImportService(repository=Mock())
        with patch("fin_ops_platform.services.tax_certified_import_service.MAX_IMPORT_RECORDS", 1):
            with self.assertRaisesRegex(ValueError, "每次认证导入"):
                service.preview_files(imported_by="owner", uploads=[certified_upload(), certified_upload(digital="SECOND")])
        service._repository.save_session.assert_not_called()
        workbook = MagicMock()
        workbook.sheetnames = ["发票"]
        consumed = 0
        def empty_rows():
            nonlocal consumed
            yield tuple(SOURCE_COLUMNS)
            for _ in range(50_100):
                consumed += 1
                yield ()
        workbook.__getitem__.return_value.iter_rows.return_value = empty_rows()
        with patch("fin_ops_platform.services.tax_certified_import_service.load_workbook", return_value=workbook):
            with self.assertRaisesRegex(ValueError, "50000"):
                self.preview()
        self.assertEqual(consumed, 50_000)
        workbook.close.assert_called_once()


class TaxCertifiedImportSummaryTests(unittest.TestCase):
    def preview(self, uploads, classifications=None, **kwargs):
        repository = Mock()
        repository.classify_rows.side_effect = lambda rows: {
            row["unique_key"]: {"blocking": False, "match_status": "outside_invoices",
                "matched_invoice_id": None, "expected_version": 0, "dedupe_status": "new",
                "error_message": None, "correction_changes": [], **(classifications or {}).get(row["unique_key"], {})} for row in rows}
        service = TaxCertifiedImportService(repository=repository)
        return TaxCertifiedImportApplicationService(certified_import_service=service).preview_payload(
            imported_by="owner", uploads=uploads, **kwargs)

    def test_mixed_file_counts_source_and_omits_normal_rows_and_ignored_metadata(self):
        result = self.preview([certified_upload(), certified_upload(label="通行费发票")])
        self.assertEqual(result["summary"], {"source_count": 2, "recognized_count": 1,
            "invalid_count": 0, "ignored_count": 1, "matched_invoice_count": 0,
            "outside_invoices_count": 1, "conflict_count": 0, "duplicate_count": 0,
            "blocking_count": 0, "new_count": 1, "relink_count": 0})
        self.assertTrue(all(file["rows"] == [] and file["missing_metadata"] == [] for file in result["files"]))

    def test_repeated_source_count_matches_commit_and_each_processing_action_is_distinct(self):
        cases = {"new": (1, 0, 1, 0), "duplicate": (0, 0, 2, 0),
                 "relink": (0, 1, 1, 0), "conflict": (0, 0, 1, 1)}
        for status, expected in cases.items():
            with self.subTest(status=status):
                result = self.preview([certified_upload(), certified_upload()],
                    {"digital:TEST-DIGITAL-1": {"dedupe_status": status}})
                counts = result["summary"]
                self.assertEqual((counts["source_count"], counts["recognized_count"]), (2, 2))
                self.assertEqual(tuple(counts[key] for key in ("new_count", "relink_count", "duplicate_count", "conflict_count")), expected)
                exceptions = [row for file in result["files"] for row in file["rows"]]
                self.assertEqual(len(exceptions), int(status == "conflict"))
                self.assertTrue(all("source_fields" not in row for row in exceptions))

    def test_missing_metadata_only_considers_eligible_rows_and_explicit_values_resolve_it(self):
        uploads = [certified_upload(month=None, buyer=None), certified_upload(label="通行费发票", month=None, buyer=None)]
        result = self.preview(uploads)
        self.assertEqual(result["files"][0]["missing_metadata"], ["month", "buyer_tax_no"])
        self.assertEqual(result["files"][1]["missing_metadata"], [])
        result = self.preview(uploads, month="2026-09", buyer_tax_no="TEST-BUYER")
        self.assertTrue(all(file["missing_metadata"] == [] for file in result["files"]))

    def test_invalid_and_blocking_rows_remain_visible_without_source_field_payload(self):
        result = self.preview([certified_upload(amount="NaN"), certified_upload(digital="BLOCKED")],
            {"digital:BLOCKED": {"blocking": True, "error_message": "发票身份匹配不唯一。"}})
        self.assertEqual(result["summary"]["invalid_count"], 2)
        self.assertEqual(result["summary"]["blocking_count"], 1)
        self.assertEqual(result["summary"]["recognized_count"], 0)
        for file in result["files"]:
            self.assertEqual(file["rows"][0]["row_status"], "invalid")
            self.assertTrue(file["rows"][0]["error_message"])
            self.assertNotIn("source_fields", file["rows"][0])
