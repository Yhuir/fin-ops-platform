from __future__ import annotations

import unittest
from decimal import Decimal
from typing import Any

from fin_ops_platform.domain.enums import InvoiceType, TransactionDirection
from fin_ops_platform.domain.models import BankTransaction, Counterparty, Invoice
from fin_ops_platform.services.imports import ImportNormalizationService
from fin_ops_platform.services.output_invoice_collection_canonical_query_service import (
    OutputInvoiceCollectionCanonicalQueryService,
)
from fin_ops_platform.services.output_invoice_collection_service import (
    OutputInvoiceCollectionError,
    OutputInvoiceCollectionQueryService,
)
from fin_ops_platform.services.workbench_relation_modes import (
    OUTPUT_INVOICE_REVERSAL_RELATION_MODE,
)


class FakeOutputCanonicalRelationReader:
    def __init__(self, relations: list[dict[str, Any]]) -> None:
        self._relations = relations

    def active_relations_for_row_ids(
        self,
        row_ids: list[str],
    ) -> list[dict[str, Any]]:
        wanted = set(row_ids)
        return [
            {**dict(relation), "status": "active"}
            for relation in self._relations
            if wanted.intersection(relation["row_ids"])
        ]


class OutputInvoiceCollectionQueryServiceTests(unittest.TestCase):
    def test_missing_source_gross_keeps_income_relation_and_unknown_pending_amount(self):
        invoice = self._invoice("missing-gross", "MISSING")
        invoice.total_with_tax = None
        bank = self._bank("source-income", "40", TransactionDirection.INFLOW)
        service = self._service(invoices=[invoice], transactions=[bank], relations=[
            self._relation("source-case", [invoice.id, bank.id], ["invoice", "bank"])])
        payload = service.list_rows()
        row = payload["rows"][0]
        self.assertEqual(row["invoice"]["totalWithTax"], "")
        self.assertEqual(row["bankTransactions"]["relationCount"], 1)
        self.assertEqual(row["collectionStatus"]["code"], "pending_collection")
        self.assertEqual(row["collectionStatus"]["collectedAmount"], "40.00")
        self.assertEqual(row["collectionStatus"]["pendingAmount"], "")
        self.assertIn("原件未提供价税合计", row["collectionStatus"]["reason"])
        self.assertEqual(payload["summary"]["pendingAmount"], "")
        self.assertEqual(payload["summary"]["collectedAmount"], "40.00")
        for field in ("total_with_tax", "pending_amount"):
            self.assertEqual(service.list_rows(filters=[{"field": field, "operator": "equals", "value": "0"}])["pagination"]["total"], 0)
            self.assertEqual(service.list_rows(filters=[{"field": field, "operator": "between", "value": {"min": "0", "max": "100"}}])["pagination"]["total"], 0)
        invoice.total_with_tax = Decimal("100")
        invoice.amount = None
        row = service.list_rows()["rows"][0]
        self.assertEqual(row["collectionStatus"]["code"], "partial_collected")
        self.assertEqual(row["collectionStatus"]["pendingAmount"], "60.00")

    def test_missing_gross_keeps_explicit_red_blue_relation_but_not_inferred_amount(self):
        blue = self._invoice("source-blue", "26532000000809302711")
        red = self._invoice("source-red", "26532000000809302712", amount="-94.34", tax_amount="-5.66",
                            total_with_tax="-100", is_positive_invoice="否",
                            remark="被红冲蓝字数电发票号码：26532000000809302711")
        blue.total_with_tax = red.total_with_tax = None
        rows = self._service(invoices=[blue, red]).list_rows()["rows"]
        self.assertEqual({row["collectionStatus"]["code"] for row in rows}, {"reversed_by_red", "reverses_blue"})
        for row in rows:
            self.assertEqual(row["invoiceRelations"]["relationCount"], 2)
            self.assertTrue(all(item["totalWithTax"] == "" for item in row["invoiceRelations"]["summaries"]))

    def test_source_zero_gross_remains_zero(self):
        invoice = self._invoice("zero-gross", "ZERO", amount="0", tax_amount="0", total_with_tax="0")
        payload = self._service(invoices=[invoice]).list_rows()
        self.assertEqual(payload["rows"][0]["invoice"]["totalWithTax"], "0.00")
        self.assertEqual(payload["rows"][0]["collectionStatus"]["pendingAmount"], "0.00")
        self.assertEqual(payload["summary"]["pendingAmount"], "0.00")

    def test_source_rates_and_unknown_facets_ignore_self_without_broadening_old_filters(self):
        rates = ["0.13", "13%", None, "0", "免税", "不征税", "mixed", "0.06"]
        invoices = [self._invoice(f"rate-{index}", str(index)) for index in range(len(rates))]
        for invoice, rate in zip(invoices, rates, strict=True):
            invoice.tax_rate = rate
        service = self._service(invoices=invoices)
        filters = [{"field": "tax_rate", "operator": "in", "values": ["0.13", "13%"]}]
        page = service.list_rows(page_size=1, filters=filters)
        self.assertEqual(page["pagination"]["total"], 2)
        self.assertEqual(page["summary"]["totalWithTax"], "200.00")
        self.assertEqual(page["summary"]["amountWithoutTax"], "188.68")
        self.assertEqual(page["appliedFilters"]["filters"][0]["values"], ["13%"])
        options = next(field["options"] for field in service.filter_options(filters=filters)["fields"] if field["field"] == "tax_rate")
        self.assertEqual({option["value"]: option["count"] for option in options},
                         {"13%": 2, "—": 1, "多税率": 1, "0%": 1, "免税": 1, "不征税": 1, "6%": 1})
        self.assertEqual(service.list_rows(filters=[{"field": "tax_rate", "operator": "in", "values": ["6%（推算）"]}])["pagination"]["total"], 0)
        for invalid in ("13%", [13], []):
            with self.subTest(invalid=invalid), self.assertRaises(OutputInvoiceCollectionError):
                service.list_rows(filters=[{"field": "tax_rate", "operator": "in", "values": invalid}])

    def test_rows_expose_only_canonical_invoice_status_bank_and_invoice_relations(self) -> None:
        row = self._service(
            invoices=[self._invoice("blue", "1001", total_with_tax="100.00")]
        ).list_rows()["rows"][0]

        self.assertEqual(
            set(row),
            {
                "id",
                "invoiceId",
                "invoiceIdentityKey",
                "invoice",
                "collectionStatus",
                "bankTransactions",
                "invoiceRelations",
            },
        )
        self.assertEqual(row["collectionStatus"]["code"], "pending_collection")
        self.assertEqual(row["collectionStatus"]["label"], "待收款")
        self.assertEqual(row["collectionStatus"]["collectedAmount"], "0.00")
        self.assertEqual(row["collectionStatus"]["pendingAmount"], "100.00")
        self.assertNotIn("oa", row)
        self.assertNotIn("receipt", row)
        self.assertNotIn("redInvoiceRelation", row)

    def test_collection_status_counts_only_linked_income_transactions(self) -> None:
        invoices = [
            self._invoice("paid", "1001", total_with_tax="100.00"),
            self._invoice("partial", "1002", total_with_tax="100.00"),
            self._invoice("pending", "1003", total_with_tax="100.00"),
        ]
        banks = [
            self._bank("income-paid", "100.00", TransactionDirection.INFLOW),
            self._bank("income-partial", "40.00", TransactionDirection.INFLOW),
            self._bank("outflow", "100.00", TransactionDirection.OUTFLOW),
        ]
        service = self._service(
            invoices=invoices,
            transactions=banks,
            relations=[
                self._relation("case-paid", ["paid", "income-paid"], ["invoice", "bank"]),
                self._relation(
                    "case-partial",
                    ["partial", "income-partial"],
                    ["invoice", "bank"],
                    matched=False,
                ),
                self._relation("case-outflow", ["pending", "outflow"], ["invoice", "bank"]),
            ],
        )

        rows = {row["invoiceId"]: row for row in service.list_rows()["rows"]}

        self.assertEqual(rows["paid"]["collectionStatus"]["code"], "collected")
        self.assertEqual(rows["partial"]["collectionStatus"]["code"], "partial_collected")
        self.assertEqual(rows["partial"]["collectionStatus"]["collectedAmount"], "40.00")
        self.assertEqual(rows["partial"]["collectionStatus"]["pendingAmount"], "60.00")
        self.assertEqual(rows["pending"]["collectionStatus"]["code"], "pending_collection")
        self.assertEqual(rows["pending"]["bankTransactions"]["receivedTotal"], "0.00")

    def test_exact_reversal_remark_drives_blue_red_and_unmatched_red_statuses(self) -> None:
        blue_invoice_no = "26532000000809302711"
        blue = self._invoice("blue", blue_invoice_no, total_with_tax="100.00")
        red = self._invoice(
            "red",
            "26532000000808367761",
            amount="-94.34",
            tax_amount="-5.66",
            total_with_tax="-100.00",
            is_positive_invoice="否",
            invoice_date="2026-05-21",
            remark=f"被红冲蓝字数电发票号码：{blue_invoice_no}",
        )
        unmatched_red = self._invoice(
            "red-unmatched",
            "2003",
            amount="-47.17",
            tax_amount="-2.83",
            total_with_tax="-50.00",
            is_positive_invoice="否",
            invoice_date="2026-05-22",
        )
        service = self._service(invoices=[blue, red, unmatched_red])

        rows = {row["invoiceId"]: row for row in service.list_rows()["rows"]}

        self.assertEqual(rows["blue"]["collectionStatus"]["code"], "reversed_by_red")
        self.assertEqual(rows["red"]["collectionStatus"]["code"], "reverses_blue")
        self.assertEqual(
            rows["red-unmatched"]["collectionStatus"]["code"],
            "unmatched_red",
        )
        self.assertEqual(
            {item["invoiceId"] for item in rows["blue"]["invoiceRelations"]["summaries"]},
            {"blue", "red"},
        )
        self.assertTrue(all(
            item["relationMode"] == OUTPUT_INVOICE_REVERSAL_RELATION_MODE
            for item in rows["blue"]["invoiceRelations"]["summaries"]
        ))

    def test_relation_case_does_not_merge_invoices_or_duplicate_bank_collection(self) -> None:
        target_no = "26532000000809302711"
        invoices = [
            self._invoice("collected-blue", "26532000000809764126", total_with_tax="182400.00", amount="172075.47", tax_amount="10324.53"),
            self._invoice("target-blue", target_no, total_with_tax="182400.00", amount="172075.47", tax_amount="10324.53"),
            self._invoice(
                "red",
                "26532000000808367761",
                total_with_tax="-182400.00",
                amount="-172075.47",
                tax_amount="-10324.53",
                is_positive_invoice="否",
                remark=f"被红冲蓝字数电发票号码：{target_no}",
            ),
        ]
        bank = self._bank("bank", "182400.00", TransactionDirection.INFLOW)
        service = self._service(
            invoices=invoices,
            transactions=[bank],
            relations=[
                self._relation(
                    "collection-case",
                    ["collected-blue", "target-blue", "red", "bank"],
                    ["invoice", "invoice", "invoice", "bank"],
                )
            ],
        )

        rows = {row["invoiceId"]: row for row in service.list_rows()["rows"]}

        self.assertEqual(set(rows), {"collected-blue", "target-blue", "red"})
        self.assertEqual(rows["collected-blue"]["collectionStatus"]["code"], "collected")
        self.assertEqual(rows["target-blue"]["collectionStatus"]["code"], "reversed_by_red")
        self.assertEqual(rows["red"]["collectionStatus"]["code"], "reverses_blue")
        self.assertEqual(rows["collected-blue"]["bankTransactions"]["receivedTotal"], "182400.00")
        self.assertEqual(rows["target-blue"]["bankTransactions"]["receivedTotal"], "0.00")
        self.assertEqual(rows["red"]["bankTransactions"]["receivedTotal"], "0.00")
        payload = service.list_rows(page=1, page_size=1)
        self.assertEqual(payload["pagination"]["total"], 3)
        self.assertEqual(payload["summary"]["invoiceCount"], 3)
        status_field = next(field for field in service.filter_options()["fields"] if field["field"] == "collection_status")
        counts = {item["value"]: item["count"] for item in status_field["options"]}
        self.assertEqual(sum(counts.values()), 3)
        self.assertEqual(counts["collected"], 1)
        self.assertEqual(counts["reversed_by_red"], 1)
        self.assertEqual(counts["reverses_blue"], 1)
        labels = {row["collectionStatus"]["label"] for row in rows.values()}
        self.assertEqual(labels, {"已收款", "已被冲", "已关联蓝字"})


    def test_filter_sort_paging_and_export_use_current_contract(self) -> None:
        service = self._service(
            invoices=[
                self._invoice("a", "3001", buyer_name="甲客户", total_with_tax="30.00"),
                self._invoice("b", "3002", buyer_name="乙客户", total_with_tax="10.00"),
                self._invoice("c", "3003", buyer_name="甲客户", total_with_tax="20.00"),
            ]
        )

        payload = service.list_rows(
            page=1,
            page_size=1,
            filters='[{"field":"buyer_name","operator":"in","values":["甲客户"]}]',
            sort_field="total_with_tax",
            sort_direction="desc",
        )
        options = service.filter_options()
        from io import BytesIO

        from openpyxl import load_workbook
        _, content = service.export()
        sheet = load_workbook(BytesIO(content)).active
        columns = [cell.value for cell in sheet[1]]

        self.assertEqual(payload["pagination"], {"page": 1, "pageSize": 1, "total": 2})
        self.assertEqual(payload["rows"][0]["invoiceId"], "a")
        fields = {field["field"] for field in options["fields"]}
        self.assertIn("collection_status", fields)
        self.assertNotIn("receipt_status", fields)
        self.assertNotIn("oa_status", fields)
        self.assertIn("价税合计", columns)
        self.assertFalse(any("收据" in column or "OA" in column for column in columns))

    def test_red_invoice_remark_drives_display_search_detail_and_export_evidence(self) -> None:
        target_invoice_no = "26532000000395506981"
        remark = f"被红冲蓝字数电发票号码：{target_invoice_no}"
        service = self._service(
            invoices=[
                self._invoice(
                    "red-with-remark",
                    "5001",
                    amount="-94.34",
                    tax_amount="-5.66",
                    total_with_tax="-100.00",
                    is_positive_invoice="否",
                    remark=remark,
                )
            ]
        )

        row = service.list_rows()["rows"][0]
        detail = service.invoice_detail("red-with-remark")
        searched = service.list_rows(keyword=target_invoice_no)
        from io import BytesIO

        from openpyxl import load_workbook
        _, content = service.export()
        sheet = load_workbook(BytesIO(content)).active
        columns = [cell.value for cell in sheet[1]]

        self.assertEqual(
            row["invoice"]["reversalTargetInvoiceNos"],
            [target_invoice_no],
        )
        self.assertEqual(detail["remark"], remark)
        self.assertEqual(
            detail["reversalTargetInvoiceNos"],
            [target_invoice_no],
        )
        self.assertEqual(searched["pagination"]["total"], 1)
        self.assertIn("备注", columns)
        self.assertEqual(
            sheet.cell(2, columns.index("备注") + 1).value,
            remark,
        )

    def test_non_contract_remark_does_not_invent_a_reversal_target(self) -> None:
        service = self._service(
            invoices=[
                self._invoice(
                    "red-with-free-text",
                    "5002",
                    is_positive_invoice="否",
                    remark="参考蓝票26532000000395506981",
                )
            ]
        )

        row = service.list_rows()["rows"][0]

        self.assertEqual(row["invoice"]["reversalTargetInvoiceNos"], [])

    def test_red_invoice_with_two_different_exact_targets_stays_unmatched(self) -> None:
        first_target = "26532000000395506981"
        second_target = "26532000000809302711"
        service = self._service(
            invoices=[
                self._invoice("blue-one", first_target),
                self._invoice("blue-two", second_target),
                self._invoice(
                    "ambiguous-red",
                    "26532000000808367761",
                    amount="-94.34",
                    tax_amount="-5.66",
                    total_with_tax="-100.00",
                    is_positive_invoice="否",
                    remark=(
                        f"被红冲蓝字数电发票号码：{first_target}；"
                        f"被红冲蓝字数电发票号码：{second_target}"
                    ),
                ),
            ]
        )

        rows = {row["invoiceId"]: row for row in service.list_rows()["rows"]}

        self.assertEqual(
            rows["ambiguous-red"]["collectionStatus"]["code"],
            "unmatched_red",
        )
        self.assertEqual(rows["ambiguous-red"]["invoiceRelations"]["summaries"], [])

    def test_inline_sources_are_projected_from_the_authorized_snapshot(self) -> None:
        invoice = self._invoice("invoice", "4001", total_with_tax="100.00")
        bank = self._bank("bank", "100.00", TransactionDirection.INFLOW)
        assembler = self._service(
            invoices=[invoice],
            transactions=[bank],
            relations=[self._relation("case", ["invoice", "bank"], ["invoice", "bank"])],
        )
        from types import SimpleNamespace
        from unittest.mock import Mock
        snapshot = SimpleNamespace(
            groups=assembler._invoice_groups(month=None, context=assembler._query_context()),
            supporting_groups=[], transactions=[bank], oa_records=[], bank_account_mappings=[], bank_labels={bank.id: []},
            relations=[{**self._relation("case", ["invoice", "bank"], ["invoice", "bank"]), "status": "active"}],
        )
        canonical = OutputInvoiceCollectionCanonicalQueryService(
            repository=Mock(load_row=Mock(return_value=snapshot)),
            row_assembler=assembler,
        )
        row_id = assembler.list_rows()["rows"][0]["id"]

        row = canonical.row_by_id(row_id)
        sources = {column["kind"]: column for column in row["relationSources"]}
        self.assertEqual(sources["bank"]["count"], 1)
        self.assertEqual(sources["bank"]["members"][0]["id"], "bank")
        self.assertEqual(sources["invoice"]["count"], 1)
        self.assertEqual(sources["invoice"]["members"][0]["id"], "invoice")
        self.assertEqual(sources["oa"]["members"], [])
        self.assertEqual(sources["bank"]["members"][0]["relationIds"], ["case"])

    def test_page_size_is_bounded(self) -> None:
        service = self._service(invoices=[])

        with self.assertRaises(OutputInvoiceCollectionError) as context:
            service.list_rows(page_size=201)

        self.assertEqual(context.exception.error_code, "invalid_paging")

    @classmethod
    def _service(
        cls,
        *,
        invoices: list[Invoice],
        transactions: list[BankTransaction] | None = None,
        relations: list[dict[str, Any]] | None = None,
    ) -> OutputInvoiceCollectionQueryService:
        return OutputInvoiceCollectionQueryService(
            import_service=ImportNormalizationService(
                existing_invoices=invoices,
                existing_transactions=transactions or [],
            ),
            relation_reader=FakeOutputCanonicalRelationReader(relations or []),
        )

    @staticmethod
    def _relation(
        case_id: str,
        row_ids: list[str],
        row_types: list[str],
        *,
        matched: bool = True,
        relation_mode: str = "manual_confirmed",
    ) -> dict[str, Any]:
        return {
            "case_id": case_id,
            "row_ids": row_ids,
            "row_types": row_types,
            "relation_mode": relation_mode,
            "amount_check": {"matched": matched},
        }

    @staticmethod
    def _invoice(
        invoice_id: str,
        invoice_no: str,
        *,
        buyer_name: str = "测试客户",
        amount: str = "94.34",
        tax_amount: str = "5.66",
        total_with_tax: str = "100.00",
        invoice_date: str = "2026-05-20",
        is_positive_invoice: str = "是",
        remark: str = "",
    ) -> Invoice:
        buyer = Counterparty(
            id=f"buyer-{invoice_id}",
            name=buyer_name,
            normalized_name=buyer_name,
            counterparty_type="customer",
            tax_no="91530000BUYER",
        )
        return Invoice(
            id=invoice_id,
            invoice_type=InvoiceType.OUTPUT,
            invoice_no=invoice_no,
            counterparty=buyer,
            amount=Decimal(amount),
            signed_amount=Decimal(total_with_tax),
            invoice_date=invoice_date,
            seller_name="云南溯源科技有限公司",
            buyer_name=buyer_name,
            seller_tax_no="91530000SELLER",
            buyer_tax_no=buyer.tax_no,
            tax_rate="6%",
            tax_amount=Decimal(tax_amount),
            total_with_tax=Decimal(total_with_tax),
            taxable_item_name="服务费",
            is_positive_invoice=is_positive_invoice,
            remark=remark,
        )

    @staticmethod
    def _bank(
        transaction_id: str,
        amount: str,
        direction: TransactionDirection,
    ) -> BankTransaction:
        return BankTransaction(
            id=transaction_id,
            account_no="622200001234",
            txn_direction=direction,
            counterparty_name_raw="测试客户",
            amount=Decimal(amount),
            signed_amount=(
                Decimal(amount)
                if direction == TransactionDirection.INFLOW
                else -Decimal(amount)
            ),
            txn_date="2026-05-21",
            trade_time="2026-05-21 10:00:00",
            imported_bank_name="建设银行",
            imported_bank_last4="1234",
            summary="服务费",
        )
