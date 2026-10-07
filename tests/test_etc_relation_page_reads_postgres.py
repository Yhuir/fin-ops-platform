from __future__ import annotations

import unittest
from dataclasses import replace

from fin_ops_platform.services.imports import ImportNormalizationService
from fin_ops_platform.services.input_invoice_usage_canonical_query_service import InputInvoiceUsageCanonicalQueryService
from fin_ops_platform.services.input_invoice_usage_service import InputInvoiceUsageQueryService
from fin_ops_platform.services.oa_payment_status_service import OAPaymentStatusRecord
from fin_ops_platform.services.oa_pending_payment_query_service import OaPendingPaymentQueryService
from fin_ops_platform.services.pending_invoice_canonical_query import (
    PendingInvoiceCanonicalQueryService,
    PostgresPendingInvoiceCanonicalRepository,
)
from fin_ops_platform.services.postgres_connection import PostgresConnection, PostgresSettings
from fin_ops_platform.services.postgres_repositories.invoice_usage_collection_query import (
    PostgresInputInvoiceUsageQueryRepository,
)
from fin_ops_platform.services.postgres_repositories.oa_pending_payment_query import (
    PostgresOaPendingPaymentQueryRepository,
)
from fin_ops_platform.services.postgres_repositories.oa_pending_payment_source_snapshot import (
    PostgresOaPendingPaymentSourceSnapshotRepository,
)
from fin_ops_platform.services.postgres_repositories.workbench_relation import PostgresWorkbenchRelationRepository
from fin_ops_platform.services.workbench_relation_command_service import WorkbenchRelationCommandService

from tests.postgres_test_utils import apply_test_migrations, require_postgres_test_database_url, truncate_test_database
from tests.test_invoice_usage_collection_postgres_integration import _UnexpectedPaymentRulesProvider
from tests.test_oa_pending_payment_postgres_integration import _in_progress_record


class EtcRelationPageReadsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.database_url = require_postgres_test_database_url()
        apply_test_migrations(cls.database_url)

    def setUp(self):
        truncate_test_database(self.database_url)
        self.connection = PostgresConnection(PostgresSettings(database_url=self.database_url, pool_enabled=False))
        PostgresOaPendingPaymentSourceSnapshotRepository(
            self.connection, relation_command_service_for_transaction=lambda tx: WorkbenchRelationCommandService(
                relation_repository=PostgresWorkbenchRelationRepository(tx)),
        ).commit_authoritative_snapshot(scope_key='2026-05', tenant_id='default', projection_records=[],
            admission_records=[replace(_in_progress_record(), amount='47.00')], payment_statuses={'flow-in-progress-1':OAPaymentStatusRecord(flow_id='flow-in-progress-1',pay_status=0)})
        self.connection.execute("""insert into app.bank_transactions(
            legacy_mongo_id, account_no, txn_direction, counterparty_name_raw, amount, signed_amount,
            txn_date, txn_month, status) values ('etc-bank', '8106', 'outflow', 'ETC还款', 47, -47,
            '2026-05-20','2026-05-01','pending')""")
        self.connection.execute("""insert into app.etc_business_batches(business_batch_id,status,scope_month,
            invoice_count,total_amount,raw_payload) values ('business-47','oa_submitted','2026-05-01',47,47,
            '{"normalized_payload":{"external_etc_batch_id":"external-47"}}')""")
        self.connection.execute("""insert into app.etc_invoices(etc_invoice_id,business_batch_id,status,invoice_no,
            invoice_date,seller_name,amount,total_with_tax)
            select 'etc-'||n,'business-47','submitted','NO-'||n,'2026-05-01','高速公路',1,1 from generate_series(1,47) n""")
        self.connection.execute("""insert into app.invoices(legacy_mongo_id,invoice_type,invoice_no,invoice_date,
            invoice_month,seller_name,buyer_name,amount,signed_amount,total_with_tax,status,etc_invoice_id)
            select 'inv-'||n,'input','NO-'||n,'2026-05-01','2026-05-01','高速公路','测试公司',1,1,1,'pending','etc-'||n
            from generate_series(1,47) n""")
        self.connection.execute("""insert into app.workbench_pair_relations(case_id,relation_mode,status,row_ids,row_types)
            values ('CASE-ETC','batch_accounting','active',
            array['oa-in-progress-1','etc-bank','etc-summary-external-47','inv-1'], array['oa','bank','invoice','invoice'])""")
        self.connection.execute("""update app.workbench_pair_relations set raw_payload=jsonb_build_object('normalized_payload',
            jsonb_build_object('case_id',case_id,'status',status,'relation_mode',relation_mode,'row_ids',row_ids,'row_types',row_types))""")

    def tearDown(self):
        self.connection.close()
        truncate_test_database(self.database_url)

    def test_all_pages_resolve_same_47_members_without_multiplying_payment(self):
        oa = OaPendingPaymentQueryService(repository=PostgresOaPendingPaymentQueryRepository(self.connection))
        rows = oa.rows({'view_mode':['in_progress']}, tenant_id='default')['rows']
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['invoice']['relationCount'], 47)
        self.assertEqual(rows[0]['bankTransaction']['paidTotal'], '47.00')
        detail = oa.relation_details(rows[0]['id'], kind='invoice', tenant_id='default')
        self.assertEqual(detail["relationCount"], 47)
        pending = PendingInvoiceCanonicalQueryService(repository=PostgresPendingInvoiceCanonicalRepository(self.connection))
        row = pending.rows({'direction':['expense'],'filter':['all']})['rows'][0]
        self.assertEqual(row['input_invoices']['relation_count'], 47)
        self.assertEqual(len({section['document_id'] for section in pending.relation_detail('etc-bank',direction='expense',kind='invoice')['sections']}),47)
        usage = InputInvoiceUsageCanonicalQueryService(repository=PostgresInputInvoiceUsageQueryRepository(self.connection),
            row_assembler=InputInvoiceUsageQueryService(import_service=ImportNormalizationService(), payment_rules_provider=_UnexpectedPaymentRulesProvider()))
        payload = usage.rows({'keyword':['NO-47']})
        self.assertEqual(len(payload['rows']), 1)
        self.assertEqual(payload['rows'][0]['oa']['relationCount'], 1)
        self.assertEqual(payload['rows'][0]['bankTransactions']['relationCount'], 1)
        self.assertEqual(self.connection.fetch_one("select row_ids from app.workbench_pair_relations where case_id='CASE-ETC'")['row_ids'],
            ['oa-in-progress-1','etc-bank','etc-summary-external-47','inv-1'])

    def test_deleted_invoice_and_withdrawn_relation_are_not_reintroduced(self):
        self.connection.execute("update app.invoices set status='deleted' where legacy_mongo_id='inv-47'")
        oa = OaPendingPaymentQueryService(repository=PostgresOaPendingPaymentQueryRepository(self.connection))
        self.assertEqual(oa.rows({'view_mode':['in_progress']},tenant_id='default')['rows'][0]['invoice']['relationCount'],46)
        self.connection.execute("update app.workbench_pair_relations set status='withdrawn' where case_id='CASE-ETC'")
        self.assertEqual(oa.rows({'view_mode':['in_progress']},tenant_id='default')['rows'][0]['invoice']['relationCount'],0)

    def test_independent_batches_sharing_invoice_keep_their_own_rows_details_and_totals(self):
        from pathlib import Path
        from tempfile import TemporaryDirectory

        from fin_ops_platform.services.input_invoice_usage_payment_rules import AppSettingsInputInvoiceUsagePaymentRulesProvider
        from fin_ops_platform.services.postgres_state_store import PostgresStateStore

        self.connection.execute("update app.invoices set status='deleted' where legacy_mongo_id in (select 'inv-'||n from generate_series(37,47) n)")
        with self.connection.transaction() as tx:
            tx.execute("set local fin_ops.correction_reason='isolated shared invoice fixture'")
            tx.execute("update app.bank_transactions set amount=36, signed_amount=-36 where legacy_mongo_id='etc-bank'")
        self.connection.execute("""insert into app.bank_transactions(legacy_mongo_id,account_no,txn_direction,
            counterparty_name_raw,amount,signed_amount,txn_date,txn_month,status)
            values ('second-bank','8106','outflow','第二批',34,-34,'2026-05-21','2026-05-01','pending')""")
        self.connection.execute("""insert into app.oa_applications(row_id,applicant,amount,workflow_status,oa_source_id,form_id,status)
            values ('second-oa','第二申请人',34,'completed','test','second','completed')""")
        self.connection.execute("""insert into app.etc_business_batches(business_batch_id,status,scope_month,
            invoice_count,total_amount,raw_payload) values ('business-second','oa_submitted','2026-05-01',34,34,
            '{"normalized_payload":{"external_etc_batch_id":"external-second"}}')""")
        self.connection.execute("""insert into app.etc_invoices(etc_invoice_id,business_batch_id,status,invoice_no,
            invoice_date,seller_name,amount,total_with_tax)
            select 'etc-'||n,'business-second','submitted','NO-'||n,'2026-05-02','高速公路',1,1 from generate_series(48,80) n""")
        self.connection.execute("""insert into app.invoices(legacy_mongo_id,invoice_type,invoice_no,invoice_date,
            invoice_month,seller_name,buyer_name,amount,signed_amount,total_with_tax,status,etc_invoice_id)
            select 'inv-'||n,'input','NO-'||n,'2026-05-02','2026-05-01','高速公路','测试公司',1,1,1,'pending','etc-'||n
            from generate_series(48,80) n""")
        self.connection.execute("""insert into app.workbench_pair_relations(case_id,relation_mode,status,row_ids,row_types)
            values ('CASE-SECOND','batch_accounting','active',array['second-oa','second-bank','etc-summary-external-second','inv-1'],
                array['oa','bank','invoice','invoice'])""")
        usage = InputInvoiceUsageCanonicalQueryService(repository=PostgresInputInvoiceUsageQueryRepository(self.connection),
            row_assembler=InputInvoiceUsageQueryService(import_service=ImportNormalizationService(), payment_rules_provider=_UnexpectedPaymentRulesProvider()))
        payload = usage.rows({})
        self.assertEqual(payload['pagination']['total'], 2)
        self.assertEqual(payload['summary']['invoiceCount'], 69)
        self.assertEqual(payload['summary']['totalWithTax'], '69.00')
        self.assertEqual(payload['classification']['all']['count'], 69)
        self.assertEqual(payload['classification']['groups'][0]['count'], 69)
        rows = {row['relationGroupId']: row for row in payload['rows']}
        for case, count in [('CASE-ETC', 36), ('CASE-SECOND', 34)]:
            row = rows[case]
            self.assertEqual(row['invoiceRelations']['relationCount'], count)
            self.assertEqual(row['oa']['relationCount'], 1)
            self.assertEqual(row['bankTransactions']['relationCount'], 1)
            self.assertEqual(row['paymentStatus']['code'], 'paid')
            for kind, expected in [('invoice', count), ('oa', 1), ('bank', 1)]:
                detail = usage.relation_details(row['id'], {'kind': [kind]})
                self.assertEqual(detail['relationCount'], expected)
        self.assertEqual(usage.rows({'keyword': ['NO-80']})['rows'][0]['relationGroupId'], 'CASE-SECOND')
        self.assertEqual(len(usage.export_rows(keyword='NO-80', limit=20000)['rows']), 34)
        self.assertEqual(len(usage.export_rows(limit=20000)['rows']), 69)
        self.connection.execute("""update app.workbench_pair_relations
            set row_ids=array['second-oa','etc-summary-external-second','inv-1'], row_types=array['oa','invoice','invoice']
            where case_id='CASE-SECOND'""")
        with TemporaryDirectory() as directory:
            store = PostgresStateStore(data_dir=Path(directory), connection=self.connection)
            provider = AppSettingsInputInvoiceUsagePaymentRulesProvider(state_store=store, transaction_factory=self.connection.transaction)
            provider.update_payment_status_rules({"expectedVersion": 1, "idempotencyKey": "shared-category", "rules": [
                {"id": "any-bank", "statusCode": "custom_shared", "label": "规则一", "priority": 1,
                 "enabled": True, "conditions": {"hasOa": True}}
            ]}, actor_id="test")
        shared = usage.rows({})['classification']
        self.assertEqual(shared['used']['count'], 69)
        self.assertEqual([group['count'] for group in shared['groups']], [36, 34])
        for parent, count, case in [('paid', 36, 'CASE-ETC'), ('unpaid', 34, 'CASE-SECOND')]:
            selected = usage.list_rows(filters=[
                {'field': 'usage_status', 'operator': 'in', 'values': ['used']},
                {'field': 'payment_group', 'operator': 'in', 'values': [parent]},
                {'field': 'payment_status', 'operator': 'in', 'values': ['custom_shared']},
            ])
            self.assertEqual(selected['summary']['invoiceCount'], count)
            self.assertEqual([row['relationGroupId'] for row in selected['rows']], [case])
            self.assertEqual(selected['classification'], shared)
        self.connection.execute("update app.workbench_pair_relations set status='withdrawn' where case_id='CASE-SECOND'")
        after = usage.rows({})
        self.assertEqual(after['classification']['used']['count'], 36)
        self.assertEqual(after['classification']['unused']['count'], 33)
        self.assertEqual(next(row for row in after['rows'] if row.get('relationGroupId') == 'CASE-ETC')['invoiceRelations']['relationCount'], 36)

    def test_export_expands_real_invoice_members_and_count_does_not_hydrate_preview(self):
        from io import BytesIO

        from fin_ops_platform.services.input_invoice_usage_export_service import InputInvoiceUsageExportService
        from openpyxl import load_workbook
        usage = InputInvoiceUsageCanonicalQueryService(repository=PostgresInputInvoiceUsageQueryRepository(self.connection),
            row_assembler=InputInvoiceUsageQueryService(import_service=ImportNormalizationService(), payment_rules_provider=_UnexpectedPaymentRulesProvider()))
        service = InputInvoiceUsageExportService(row_export_loader=usage.export_rows)
        summary = service.export_summary()
        self.assertEqual(summary["row_count"], 47)
        self.assertNotIn("sample_rows", summary)
        self.assertEqual(usage.export_rows(limit=0)["rows"], [])
        _, content = service.export()
        sheet = load_workbook(BytesIO(content)).active
        self.assertEqual(sheet.max_row, 48)
        self.assertEqual({sheet.cell(i, 2).value for i in range(2,49)}, {f"NO-{n}" for n in range(1,48)})
        self.assertEqual(sum(sheet.cell(i,12).value for i in range(2,49)),47)
        self.assertFalse(any("ID" in str(cell.value) for cell in sheet[1]))
        self.connection.execute("update app.invoices set status='deleted' where legacy_mongo_id='inv-47'")
        self.assertEqual(service.export_summary()["row_count"],46)

    def test_pending_export_expands_banks_without_duplicating_group_amounts(self):
        from decimal import Decimal
        from io import BytesIO

        from fin_ops_platform.services.pending_invoice_service import PendingInvoiceQueryService
        from openpyxl import load_workbook
        self.connection.execute("""insert into app.bank_transactions(legacy_mongo_id,account_no,txn_direction,counterparty_name_raw,amount,signed_amount,txn_date,txn_month,status)
            values ('etc-bank-2','8106','outflow','ETC还款',24,-24,'2026-05-20','2026-05-01','pending')""")
        self.connection.execute("""update app.workbench_pair_relations set row_ids=array_append(row_ids,'etc-bank-2'),row_types=array_append(row_types,'bank')""")
        query = PendingInvoiceCanonicalQueryService(repository=PostgresPendingInvoiceCanonicalRepository(self.connection))
        self.assertEqual(query.export_summary({})["row_count"],2)
        rows = query.all_rows({})["rows"]
        self.assertEqual(len(rows),2)
        exported = [PendingInvoiceQueryService._export_row(i,row) for i,row in enumerate(rows,1)]
        self.assertEqual(sorted(row["借方金额"] for row in exported),[Decimal('24'),Decimal('47')])
        self.assertEqual(len({row['bank_transactions']['primary']['bank_transaction_id'] for row in rows}),2)
