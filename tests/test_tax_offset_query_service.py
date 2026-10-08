from __future__ import annotations

import unittest
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from fin_ops_platform.services.tax_offset_export_service import TaxOffsetExportError, TaxOffsetExportService
from fin_ops_platform.services.tax_offset_query_service import DEFAULT_EXPORT_FIELDS, TaxOffsetQuery, TaxOffsetQueryService


class TaxOffsetQueryTests(unittest.TestCase):
    def test_defaults_filter_types_months_and_sorts(self):
        self.assertEqual(TaxOffsetQuery.parse({}), TaxOffsetQuery())
        query = TaxOffsetQuery.parse({"issue_month": "2026-09", "selection_month": "2026-10", "page": "2",
                                      "page_size": 100, "status": "certified", "sort_by": "selection_time", "sort_direction": "asc"})
        self.assertEqual(query.page, 2)
        for invalid in ({"month": "2026-10"}, {"issue_month": "2026-13"}, {"selection_month": "202610"},
                        {"page": True}, {"page_size": 201}, {"page": "1.1"}, {"sort_by": "id;delete"},
                        {"sort_direction": "arbitrary"}, {"status": "saved"}, {"search": ["x"]}):
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                TaxOffsetQuery.parse(invalid)

    def test_query_and_export_use_identical_filter_contract_and_catalog(self):
        calls = []
        def load(query, **kwargs):
            calls.append((query, kwargs))
            return {"rows": [], "total": 0, "page": query.page, "page_size": query.page_size, "summary": {}}
        service = TaxOffsetQueryService(canonical_repository=SimpleNamespace(load_page=load))
        filters = {"status": "uncertified", "issue_month": "2026-01", "page": 3}
        payload = service.list_payload(filters)
        service.export_rows(filters, limit=20)
        self.assertEqual(calls[0][0], calls[1][0])
        self.assertEqual(calls[1][1], {"limit_override": 20})
        self.assertEqual(len(payload["export_fields"]), 21)
        self.assertEqual({row["key"] for row in payload["export_fields"] if row["default_selected"]}, set(DEFAULT_EXPORT_FIELDS))


class TaxOffsetExportTests(unittest.TestCase):
    def test_default_eight_columns_null_zero_long_id_and_formula_are_preserved(self):
        row = {"sequence": 1, "digital_invoice_no": "0001234567890123456789", "issue_date": "2026-10-01",
               "seller_tax_no": "00000000012345", "seller_name": "=HYPERLINK(\"bad\")",
               "amount": None, "tax_amount": "0.00", "deductible_tax_amount": "1.234567"}
        calls = []
        service = TaxOffsetExportService(query_service=SimpleNamespace(
            export_rows=lambda filters, **kwargs: calls.append((filters, kwargs)) or {"rows": [row], "total": 1}))
        filename, content = service.export({"status": "certified"})
        self.assertTrue(filename.startswith("专票清单-"))
        workbook = load_workbook(BytesIO(content))
        sheet = workbook.active
        self.assertEqual(sheet.max_column, 8)
        self.assertEqual(sheet.cell(2, 2).value, row["digital_invoice_no"])
        self.assertEqual(sheet.cell(2, 5).data_type, "s")
        self.assertEqual(sheet.cell(2, 5).value, row["seller_name"])
        self.assertIsNone(sheet.cell(2, 6).value)
        self.assertEqual(sheet.cell(2, 7).value, 0)
        self.assertEqual(sheet.cell(2, 8).value, 1.234567)
        workbook.close()
        self.assertEqual(calls[0][0], {"status": "certified"})

    def test_field_selection_order_status_and_invalid_fields(self):
        service = TaxOffsetExportService(query_service=SimpleNamespace(export_rows=lambda *args, **kwargs: {
            "total": 1, "rows": [{"certification_status": "uncertified", "tax_period": None}]}))
        _, content = service.export({}, ["tax_period", "certification_status"])
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(list(workbook.active.values), [("所属期", "认证状态"), (None, "未认证")])
        workbook.close()
        for fields in ([], ["password"], ["amount", "amount"], "amount", [None]):
            with self.subTest(fields=fields), self.assertRaises(TaxOffsetExportError):
                service.export({}, fields)
        service = TaxOffsetExportService(query_service=SimpleNamespace(export_rows=lambda *args, **kwargs: {"rows": [], "total": 20001}))
        with self.assertRaisesRegex(TaxOffsetExportError, "20000"):
            service.export({})
